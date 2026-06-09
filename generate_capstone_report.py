# generate_capstone_report.py — SAFEVIEW 캡스톤 결과보고서 워드 생성
# 실행: python generate_capstone_report.py

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUTPUT = os.path.join(os.path.dirname(__file__), "SAFEVIEW_결과보고서.docx")

doc = Document()

# ── 기본 스타일 (맑은 고딕, 11pt) ──────────────────────────
style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

# 페이지 여백
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)


def set_run_font(run, size=11, bold=False, color=None, font="맑은 고딕"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font)


def add_title(text, size=22):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=RGBColor(46, 125, 50))


def add_h1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True, color=RGBColor(46, 125, 50))
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def add_h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=RGBColor(33, 33, 33))
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)


def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11)


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(2)
    run = p.runs[0] if p.runs else p.add_run("")
    run.text = text
    set_run_font(run, size=11)


def add_table_2col(rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(rows):
        c1 = table.cell(i, 0)
        c2 = table.cell(i, 1)
        c1.text = ""
        c2.text = ""
        r1 = c1.paragraphs[0].add_run(k)
        r2 = c2.paragraphs[0].add_run(v)
        set_run_font(r1, size=10, bold=True)
        set_run_font(r2, size=10)
    table.columns[0].width = Cm(4.5)
    table.columns[1].width = Cm(11.5)


# ══════════════════════════════════════════════════════
# 표지
# ══════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
add_title("캡스톤디자인 결과보고서", size=20)
doc.add_paragraph()
add_title("SAFEVIEW", size=32)
add_title("AI 기반 사각지대 위험 감지 시스템", size=16)
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

info_p = doc.add_paragraph()
info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_run = info_p.add_run(
    "팀명 : 세이프뷰(SAFEVIEW)\n"
    "프로젝트명 : SAFEVIEW — AI 기반 사각지대 위험 감지 시스템\n"
    "기간 : 2026년 3월 ~ 6월\n"
    "개발 환경 : Python · Streamlit · YOLOv8 · OpenCV · RTSP"
)
set_run_font(info_run, size=12)

doc.add_page_break()


# ══════════════════════════════════════════════════════
# 1. 연구 배경 및 목적
# ══════════════════════════════════════════════════════
add_h1("1. 연구 배경 및 목적")

add_body(
    "노상공영주차장과 도로변 주차 구간 그리고 좁은 골목길에서는 주차 차량과 구조물 "
    "때문에 운전자와 보행자의 시야가 가려지는 사각지대가 자주 형성된다. 이런 환경에서는 "
    "보행자와 차량이 서로를 늦게 인지하게 되고 인지 시점에는 이미 거리가 너무 가까워 "
    "충분한 반응 시간을 확보하기 어렵다. 그 결과 급정지나 급회피가 필요한 위험 상황으로 "
    "이어지기 쉬우며 실제로 생활도로에서 발생하는 보행자 사고의 상당수가 이런 환경적 "
    "원인과 연관되어 있다."
)
add_body(
    "현재 운영 중인 일반 CCTV는 대부분 사고가 일어난 뒤 영상을 돌려보는 사후 확인 목적으로 "
    "쓰이고 있다. 사고 자체를 줄이려면 위험 상황을 사전에 감지하고 즉시 알릴 수 있는 능동형 "
    "시스템이 필요하다. 본 연구는 이런 문제의식에서 출발해 카메라 영상에서 사람과 차량을 "
    "실시간으로 인식하고 사용자가 직접 지정한 위험 구역(ROI)에 보행자가 진입할 경우 "
    "시각과 청각 경고를 동시에 출력하는 사각지대 위험 감지 시스템을 개발하는 것을 목적으로 "
    "한다."
)
add_body(
    "본 연구의 구체적인 목적은 다음과 같다. 첫째 일반 환경에 설치된 CCTV 영상으로부터 "
    "사람과 차량을 실시간으로 인식한다. 둘째 사용자가 화면 위에 마우스로 다각형 위험 구역을 "
    "직접 설정할 수 있도록 한다. 셋째 사람·차량·ROI 침입 조건을 기반으로 위험 상황을 판정해 "
    "즉시 시청각 경고를 제공한다. 넷째 위험 발생 시점의 이미지와 영상 클립을 자동으로 "
    "저장하고 로그를 남겨 사후 분석이 가능하도록 한다."
)


# ══════════════════════════════════════════════════════
# 2. 프로젝트의 범위
# ══════════════════════════════════════════════════════
add_h1("2. 프로젝트의 범위")

add_body(
    "본 프로젝트는 차량 탑재용 자율주행 시스템이나 모든 도로 환경을 대상으로 하는 범용 "
    "교통 안전 시스템을 만드는 것이 아니다. 주차 차량으로 형성된 시야 제한 구간을 중심으로 "
    "한 소프트웨어 프로토타입 개발에 범위를 한정한다."
)

add_h2("포함되는 범위")
add_bullet("RTSP CCTV와 영상 파일 두 가지 입력 소스 대응")
add_bullet("YOLOv8 기반 사람·차량 실시간 객체 인식")
add_bullet("마우스 드로잉 방식의 다각형 ROI 설정과 카메라별 저장")
add_bullet("사람 + 차량 + ROI 침입 조건의 규칙 기반 위험 판정")
add_bullet("화면 시각 강조와 경고음을 결합한 시청각 알림")
add_bullet("위험 이벤트의 캡처 이미지 · 전 5초 + 후 10초 영상 클립 · CSV 로그 자동 저장")
add_bullet("다중 CCTV 프리셋 토글 기능")
add_bullet("정지 차량 자체 판정 알고리즘으로 의미 없는 경고 차단")
add_bullet("야간 자동 밝기 보정(CLAHE) 기능")
add_bullet("KISA 지능형 CCTV 성능시험 인증제도 기준의 정량 성능 평가 페이지")
add_bullet("Cloudflare Tunnel을 활용한 외부 시연 공유")

add_h2("제외되는 범위")
add_bullet("차량 탑재용 임베디드 하드웨어 설계 및 제작")
add_bullet("자율주행 수준의 정밀 충돌 예측")
add_bullet("악천후와 모든 시간대를 포함한 도로 환경 일반화")
add_bullet("자체 학습 데이터셋 구축 및 모델 재학습")


# ══════════════════════════════════════════════════════
# 3. 연구 가설 및 기대 효과
# ══════════════════════════════════════════════════════
add_h1("3. 연구 가설 및 기대 효과")

add_h2("3.1 연구 가설")
add_body(
    "본 연구의 가설은 다음과 같다. 사용자가 직접 설정한 ROI 영역과 사람·차량 실시간 객체 "
    "감지 결과를 결합하면 사후 확인 중심의 기존 CCTV로는 잡아내기 어려운 사각지대 위험 "
    "상황을 사전에 감지하고 경고할 수 있다. 또한 단순 객체 인식만으로는 발생하는 의미 없는 "
    "경고(주차된 차량을 위험 요소로 인식하는 사례)를 정지 차량 분리 알고리즘으로 줄일 수 "
    "있다."
)

add_h2("3.2 기대 효과")
add_body(
    "기술적 측면에서는 객체 인식 ROI 위험 판단 시청각 경고 이벤트 저장으로 이어지는 "
    "비전 기반 안전 시스템의 전체 흐름을 하나의 통합 시스템으로 제시할 수 있다. 별도의 "
    "대규모 하드웨어 없이 소프트웨어 중심으로 구현하기 때문에 기존 CCTV 인프라를 그대로 "
    "활용해 비교적 낮은 비용으로 예방형 안전 보조 체계를 구축할 수 있다는 점도 의미 있다."
)
add_body(
    "활용 측면에서는 학교 통학로 상가 밀집 지역 지자체 노상공영주차장 소규모 주차장 진출입로 "
    "골목 어귀 등 다양한 환경으로 확장이 가능하다. 사고 예방 효과뿐 아니라 자동 저장된 "
    "이벤트 영상을 통해 사고 발생 시 책임 소재 판단과 분석 자료로도 활용할 수 있다."
)


# ══════════════════════════════════════════════════════
# 4. 작품 개요
# ══════════════════════════════════════════════════════
add_h1("4. 작품 개요")

add_h2("4.1 작품의 정의 및 특징")
add_body(
    "SAFEVIEW는 일반 CCTV 영상에서 사람과 차량을 실시간으로 인식하고 사용자가 지정한 "
    "위험 구역으로 보행자가 진입할 경우 즉시 시청각 경고를 제공하는 AI 기반 사각지대 "
    "위험 감지 시스템이다. 단순한 영상 기록 장치가 아니라 위험 상황을 능동적으로 감지하고 "
    "경고하는 예방형 안전 보조 시스템이라는 점이 가장 큰 특징이다."
)
add_body("주요 특징은 다음과 같다.")
add_bullet("실시간 처리에 초점을 둔 경량 모델(YOLOv8n) 채택")
add_bullet("사용자가 화면을 보면서 마우스 클릭만으로 ROI 다각형을 설정")
add_bullet("정지 차량 자체 판정으로 주차된 차량 때문에 발생하는 오탐 차단")
add_bullet("야간이나 어두운 환경에서도 동작할 수 있도록 CLAHE 기반 자동 밝기 보정")
add_bullet("여러 대의 CCTV 정보를 프리셋으로 등록해 토글 전환")
add_bullet("위험 이벤트의 캡처 이미지와 영상 클립을 자동 저장하고 캘린더 기반 다시보기 제공")
add_bullet("KISA 인증 평가 기준에 맞춘 자체 성능 평가 페이지 내장")

add_h2("4.2 작품의 기능 및 구조")

add_body("본 시스템은 입력 → 처리 → 출력 → 저장의 네 단계로 구성되며 각 단계에 다음과 같은 모듈이 배치되어 있다.")

add_table_2col([
    ("구분",            "구성 내용"),
    ("입력 단계",       "RTSP CCTV 스트림 / 영상 파일"),
    ("처리 단계",       "OpenCV 프레임 처리, YOLOv8 객체 인식, ROI 침입 판정, 정지 차량 분리, 야간 보정(CLAHE)"),
    ("출력 단계",       "Streamlit UI 실시간 영상 표시, 화면 테두리 강조, 경고 텍스트, 경고음 재생"),
    ("저장 단계",       "캡처 이미지 / 전 5초 + 후 10초 영상 클립 / CSV 로그"),
    ("운용 편의 기능",  "다중 CCTV 프리셋 토글, 캘린더 다시보기, 외부 공유(Cloudflare Tunnel)"),
    ("평가 기능",       "KISA 기준 정탐률·오탐률·미탐률·응답시간 측정"),
])

add_body(
    "사용자 화면은 Streamlit 기반 멀티 페이지로 구성된다. 대시보드 페이지에서는 시스템 "
    "상태와 최근 이벤트를 한눈에 보여주고 모니터링 페이지에서는 실시간 영상과 경고를 "
    "확인할 수 있다. ROI 설정 페이지에서는 마우스로 위험 구역을 직접 그릴 수 있고 "
    "이벤트 다시보기 페이지에서는 날짜별로 저장된 이벤트를 캘린더 형태로 조회할 수 있다. "
    "마지막 성능 평가 페이지에서는 KISA 평가 기준에 따른 정량 지표를 입력하고 자동 "
    "계산된 결과를 확인할 수 있다."
)


# ══════════════════════════════════════════════════════
# 5. 제작 과정
# ══════════════════════════════════════════════════════
add_h1("5. 제작 과정")

add_h2("5.1 디자인 및 계획 단계")
add_body(
    "초기 단계에서는 문제 정의와 대표 시나리오 설정에 시간을 가장 많이 썼다. 단순히 "
    "‘CCTV로 사람을 인식한다’가 아니라 ‘어떤 환경의 어떤 상황을 어느 시점에 감지해야 의미가 "
    "있는가’를 구체화하는 작업이었다. 팀원과 회의를 거쳐 다음과 같이 정리했다."
)
add_bullet("타깃 환경 : 노상공영주차장 · 도로변 주차 구간 · 좁은 골목길")
add_bullet("타깃 상황 : 주차 차량 사이로 보행자가 갑자기 등장해 차량과 접근하는 시나리오")
add_bullet("감지 시점 : 사고가 발생한 뒤가 아닌 사람과 차량이 같은 위험 구역 안에 동시에 존재하는 순간")

add_body(
    "기술 스택은 ‘짧은 기간 안에 시연 가능한 통합 시스템을 만든다’는 목표에 맞춰 선정했다. "
    "Python 생태계 안에서 객체 인식 영상 처리 UI 구성을 모두 처리할 수 있도록 YOLOv8 "
    "OpenCV Streamlit 조합으로 결정했다. 또한 개발 보조 도구로 생성형 AI를 활용해 "
    "프로토타입 구현 속도를 높이되 모든 요구사항 정의 결과 검증 수정 통합은 팀원이 직접 "
    "수행하는 방식을 채택했다."
)

add_h2("5.2 구현 및 제작 단계")
add_body("구현은 다음과 같은 순서로 단계별 통합 방식으로 진행했다.")
add_bullet("1단계 : 영상 입력(RTSP / 파일) + YOLOv8 객체 인식 기본 동작 확인")
add_bullet("2단계 : 마우스 클릭 기반 다각형 ROI 설정 기능 구현 (streamlit-image-coordinates 활용)")
add_bullet("3단계 : 위험 판단 로직 — 사람 발 위치가 ROI 내부 + 차량 동시 감지 시 위험 판정")
add_bullet("4단계 : 시각 경고(화면 테두리·텍스트) + 청각 경고(경고음) 연동")
add_bullet("5단계 : 이벤트 저장 — 캡처 이미지 + 전 5초 + 후 10초 영상 클립 + CSV 로그")
add_bullet("6단계 : 다중 CCTV 프리셋 토글 기능 추가 (JSON 기반 저장)")
add_bullet("7단계 : 정지 차량 자체 판정 알고리즘 추가 (좌표 거리 + 시간 기반)")
add_bullet("8단계 : 야간 자동 밝기 보정(CLAHE) 추가")
add_bullet("9단계 : 캘린더 기반 이벤트 다시보기 페이지와 H.264 자동 변환")
add_bullet("10단계 : KISA 기준 정량 성능 평가 페이지 구현")
add_bullet("11단계 : Cloudflare Tunnel을 통한 외부 시연 환경 구성")

add_body(
    "각 단계마다 동작 확인 후 다음 단계로 넘어가는 점진적 통합 방식을 유지했기 때문에 "
    "기능 추가 과정에서 발생한 문제를 비교적 빠르게 발견하고 수정할 수 있었다."
)

add_h2("5.3 테스트 및 개선 단계")
add_body(
    "실제 위험 상황은 발생 빈도가 낮기 때문에 시연용 영상은 팀원이 직접 연출해서 촬영했다. "
    "주차 차량 사이로 보행자가 등장하는 시나리오 좁은 골목에서 차량이 진입하는 시나리오 "
    "야간 환경에서 보행자가 ROI에 진입하는 시나리오 등 여러 케이스를 촬영해 시스템 동작을 "
    "검증했다."
)
add_body("테스트 결과 발견한 주요 이슈와 개선 사항은 다음과 같다.")
add_bullet("주차된 차량이 위험 요소로 인식되는 문제 → 정지 차량 자체 판정 알고리즘 추가")
add_bullet("야간 영상에서 사람 검출률이 떨어지는 문제 → CLAHE 기반 자동 밝기 보정 추가")
add_bullet("저장된 mp4 클립이 브라우저에서 재생되지 않는 문제 → imageio-ffmpeg로 H.264 자동 변환")
add_bullet("객체 추적(ByteTrack) 도입 시 프레임이 떨어지는 문제 → 추적 기능 롤백 후 단순 검출로 복귀")
add_bullet("st.rerun() 호출 때문에 영상이 깜빡이는 문제 → st.empty() 플레이스홀더 방식으로 변경")
add_body(
    "이 외에도 외부 공유 단계에서 ngrok 무료 대역폭 초과 문제가 발생해 Cloudflare Tunnel로 "
    "전환했고 브라우저 다크 모드에서 텍스트가 잘 보이지 않는 문제는 .streamlit/config.toml "
    "설정으로 라이트 테마를 강제하는 방식으로 해결했다."
)


# ══════════════════════════════════════════════════════
# 6. 결과 분석
# ══════════════════════════════════════════════════════
add_h1("6. 결과 분석")

add_h2("6.1 작품 성능 및 기능 분석")
add_body(
    "본 시스템은 계획서에서 설정한 핵심 기능을 모두 구현 완료한 상태이며 추가로 정지 차량 "
    "분리 야간 보정 다중 CCTV 토글 성능 평가 페이지까지 확장 구현했다. 기능별 구현 결과는 "
    "다음과 같다."
)
add_table_2col([
    ("기능",                  "결과"),
    ("실시간 객체 인식",      "YOLOv8n 기준 일반 PC 환경에서 안정 동작 (FRAME_SKIP=2 적용)"),
    ("ROI 설정",              "마우스 클릭 다각형 방식 정상 동작 / 카메라 소스별 저장·재사용"),
    ("위험 판단",             "사람 + 차량 + ROI 내부 조건 기반 정상 동작"),
    ("정지 차량 분리",        "10초 이상 30px 미만 이동 시 pk.car로 분류 / 위험 대상 제외"),
    ("야간 보정",             "CLAHE 적용 후 어두운 영상에서 검출률 개선 확인"),
    ("이벤트 저장",           "캡처 이미지 + 전 5초 + 후 10초 클립 + CSV 로그 모두 정상"),
    ("다시보기",              "캘린더 기반 날짜 조회 + H.264 변환 자동 캐싱 정상"),
    ("다중 CCTV 토글",        "드롭다운 방식 프리셋 전환 정상"),
    ("외부 공유",             "Cloudflare Tunnel HTTPS 링크로 외부 시연 가능"),
    ("성능 평가 페이지",      "KISA 기준 정탐률·오탐률·미탐률·응답시간 자동 계산"),
])

add_h2("6.2 결과 비교 및 평가")
add_body(
    "기존 일반 CCTV 시스템과 본 작품을 비교하면 다음과 같은 차이를 보인다. 기존 CCTV는 "
    "녹화 중심이라 사고가 발생한 뒤에야 영상을 확인할 수 있는 반면 본 시스템은 위험 상황을 "
    "실시간으로 감지해 즉시 경고하고 동시에 사고 시점의 영상을 자동으로 저장한다. 즉 사후 "
    "확인용 도구에서 사전 예방용 도구로 활용 목적이 확장된다."
)
add_body(
    "단순 객체 인식만 적용한 시스템과 비교하면 정지 차량 자체 판정 알고리즘 덕분에 주차된 "
    "차량을 위험 요소로 인식해 발생하는 의미 없는 경고가 크게 줄었다. 또한 KISA 인증 평가 "
    "방식을 참고한 성능 평가 페이지를 내장해 정량적인 자체 검증이 가능하다는 점도 본 "
    "작품만의 특징이다."
)


# ══════════════════════════════════════════════════════
# 7. 문제점 및 해결 방안
# ══════════════════════════════════════════════════════
add_h1("7. 문제점 및 해결 방안")

add_h2("7.1 문제점 분석")
add_body(
    "프로젝트 진행 중 발생한 주요 문제점은 크게 세 가지 유형으로 나눌 수 있다. 첫 번째는 "
    "라이브러리 호환성 문제 두 번째는 성능과 정확도 사이의 트레이드오프 문제 세 번째는 "
    "외부 환경과 브라우저 환경의 차이로 인한 문제다."
)
add_bullet(
    "라이브러리 호환성 : streamlit-drawable-canvas가 최신 Streamlit과 호환되지 않아 "
    "ROI 드로잉 기능을 정상 구현하기 어려웠다."
)
add_bullet(
    "성능 vs 정확도 : 객체 추적(ByteTrack)을 도입하니 정확도는 올라갔지만 프레임이 떨어져 "
    "실시간 감시라는 핵심 가치를 해쳤다."
)
add_bullet(
    "코덱 호환성 : OpenCV의 기본 mp4v 코덱으로 저장한 영상이 브라우저에서 재생되지 않아 "
    "이벤트 다시보기 기능이 무용지물이 됐다."
)
add_bullet(
    "야간 정확도 : 어두운 영상에서는 YOLO 검출률이 눈에 띄게 떨어졌다."
)
add_bullet(
    "외부 공유 : ngrok 무료 플랜의 대역폭 제한 때문에 외부 시연이 자주 끊겼다."
)
add_bullet(
    "오탐 : 주차된 차량까지 위험 요소로 인식해 의미 없는 경고가 다수 발생했다."
)

add_h2("7.2 해결 방안 및 개선 방향")
add_body("각 문제에 대해 적용한 해결 방안은 다음과 같다.")
add_table_2col([
    ("문제",                                "해결 방안"),
    ("drawable-canvas 호환성",              "streamlit-image-coordinates로 교체 후 좌표를 직접 받아 다각형 ROI 구성"),
    ("ByteTrack 도입 시 FPS 저하",          "추적 기능 롤백 후 단순 검출 + 좌표 비교 기반 자체 정지 차량 판정 알고리즘 구현"),
    ("브라우저 재생 불가",                  "imageio-ffmpeg로 H.264 자동 변환 후 캐싱 폴더에 저장"),
    ("야간 검출률 저하",                    "CLAHE 기반 자동 밝기 보정 적용 (DARK_THRESHOLD 기준 자동 ON/OFF)"),
    ("ngrok 대역폭 초과",                   "Cloudflare Tunnel로 전환해 무료 무제한 HTTPS 공유 환경 구축"),
    ("주차 차량 오탐",                      "10초 이상 30px 미만 이동 객체를 pk.car로 분류해 위험 판정에서 제외"),
    ("st.rerun() 영상 깜빡임",              "st.empty() 플레이스홀더 + while 루프로 페이지 새로고침 없이 갱신"),
    ("브라우저 다크 모드 텍스트 가독성",     ".streamlit/config.toml에서 라이트 테마 강제 설정"),
])


# ══════════════════════════════════════════════════════
# 8. 결론
# ══════════════════════════════════════════════════════
add_h1("8. 결론")

add_h2("8.1 연구 결과 요약")
add_body(
    "본 프로젝트는 노상공영주차장과 도로변 주차 구간의 사각지대 위험에 대응하기 위한 "
    "비전 기반 안전 보조 시스템 SAFEVIEW를 개발했다. 카메라 영상에서 사람과 차량을 "
    "실시간으로 인식하고 사용자가 직접 설정한 위험 구역(ROI)에 보행자가 진입할 경우 "
    "즉시 시청각 경고를 제공하며 위험 이벤트는 캡처 이미지와 영상 클립 CSV 로그로 자동 "
    "저장된다. 계획서에서 설정한 핵심 기능 외에도 다중 CCTV 프리셋 정지 차량 분리 야간 "
    "자동 밝기 보정 KISA 기준 정량 성능 평가 기능을 추가 구현했다."
)

add_h2("8.2 연구의 의의 및 한계점")
add_body(
    "연구의 의의는 사후 확인 중심의 기존 CCTV 운용 방식을 사전 감지 중심으로 전환할 수 "
    "있는 실용 가능한 프로토타입을 제시했다는 점이다. 별도의 대규모 하드웨어 없이 "
    "소프트웨어 중심으로 구현했기 때문에 기존 CCTV 인프라를 그대로 활용할 수 있어 "
    "현실적인 확장 가능성도 확보했다."
)
add_body("다만 다음과 같은 한계점이 있다.")
add_bullet("단일 카메라 시점 기반으로 동작하므로 사각지대 전체를 완전히 커버하지는 못한다.")
add_bullet("위험 판정은 ROI 침입 여부 + 객체 동시 감지라는 규칙 기반이라 정밀 충돌 예측은 불가능하다.")
add_bullet("GPU 없는 환경을 전제로 개발했기 때문에 YOLOv8n을 사용해 정확도보다는 실시간성에 우선순위를 두었다.")
add_bullet("악천후 강한 역광 등 특수한 환경 조건에 대한 일반화는 검증되지 않았다.")

add_h2("8.3 향후 연구 방향")
add_body("본 연구를 기반으로 다음과 같은 방향으로 발전시킬 수 있다.")
add_bullet("위험 구역 자동 설정 — 차량 객체 인식 결과를 바탕으로 차량 사이 공간을 자동으로 ROI 후보로 추천")
add_bullet("물리적 경고 장치 연동 — 화면 경고와 경고음을 넘어 현장 전광판이나 스피커와 연결")
add_bullet("자체 데이터셋 기반 모델 재학습 — 사각지대 환경 영상으로 파인튜닝하여 도메인 정확도 향상")
add_bullet("다중 시점 융합 — 여러 대의 CCTV 영상을 결합해 사각지대 커버리지를 확장")
add_bullet("이벤트 분석 자동화 — 누적된 이벤트 로그를 분석해 위험 빈도가 높은 시간대와 장소를 자동 도출")

add_body(
    "본 프로젝트는 학부 캡스톤디자인 수준의 프로토타입이지만 실제 환경에서 동작 가능한 "
    "통합 시스템 형태로 완성했다는 점에서 향후 더 큰 규모의 안전 시스템으로 확장될 수 있는 "
    "기반을 마련했다."
)


# 저장
doc.save(OUTPUT)
print(f"[완료] 보고서 저장 → {OUTPUT}")

# 자동 열기
try:
    os.startfile(OUTPUT)
    print("[열기] Word에서 파일을 여는 중...")
except Exception as e:
    print(f"[실패] 파일 열기 실패: {e}")
